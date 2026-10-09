"""«Один кабинет — несколько названий»: Малый зал = «423»; «м/ф» (малое фойе) — отдельный кабинет."""
from __future__ import annotations

import json
from io import BytesIO

import openpyxl
from docx import Document
from fastapi.testclient import TestClient

from app import backup, db, room_aliases
from app import main as main_module
from app.conflicts import find_conflicts
from app.models import ConflictType


def test_defaults_when_no_file_and_mapping(tmp_path):
    aliases = room_aliases.load_aliases(str(tmp_path / "нет.json"))

    m = room_aliases.mapping(aliases)
    assert m["малый зал"] == "423"                                                     # Малый зал = 423
    assert m["мф"] == m["малое фойе"] == m["фойе малое"] == "м/ф"                     # малое фойе — свой кабинет
    assert "м/ф" not in m and m.get("423") is None                                     # м/ф и 423 НЕ сливаются
    assert room_aliases.labels(aliases) == {"423": "Малый зал", "м/ф": "Малое фойе"}
    assert room_aliases.display("423", aliases) == "Малый зал (423)"
    assert room_aliases.display("м/ф", aliases) == "Малое фойе (м/ф)" and room_aliases.display("417", aliases) == "417"


def test_apply_merges_hall_names_so_the_clash_between_423_and_maly_zal_is_found(make_lesson):
    """Один преподаватель написал «423», другой «Малый зал» — в одно время это одно помещение.
    А «м/ф» (малое фойе) — другое место, накладки с ним нет."""
    a = make_lesson(day=1, start="10:15", teacher="Иванов И.И.", student="Алексеев Ян", room="423")
    b = make_lesson(day=1, start="10:15", teacher="Петров П.П.", student="Борисова Нина", room="малый зал")
    foyer = make_lesson(day=1, start="10:15", teacher="Сидоров С.С.", student="Васильев Олег", room="м/ф")
    assert not [c for c in find_conflicts([a, b]) if c.type == ConflictType.ROOM_DOUBLE_BOOKED]   # без списка — не видно

    room_aliases.apply([a, b, foyer], room_aliases.load_aliases("/нет/файла.json"))

    assert a.room_normalized == b.room_normalized == "423" and foyer.room_normalized == "м/ф"
    clashes = [c for c in find_conflicts([a, b, foyer]) if c.type == ConflictType.ROOM_DOUBLE_BOOKED]
    assert len(clashes) == 1 and {clashes[0].lesson_a.teacher_name, clashes[0].lesson_b.teacher_name} == {"Иванов И.И.", "Петров П.П."}


def test_add_and_remove_names_and_one_name_belongs_to_one_room():
    al = {}
    assert room_aliases.add_name(al, "Малая гостиная", "305", "Гостиная")
    assert al == {"305": {"label": "Гостиная", "names": ["Малая гостиная"]}}
    assert room_aliases.add_name(al, "малая гостиная", "406")           # то же название переехало к другому кабинету
    assert al["305"]["names"] == [] and al["406"]["names"] == ["малая гостиная"]
    assert not room_aliases.add_name(al, "", "406") and not room_aliases.add_name(al, "406", "406")
    room_aliases.remove_name(al, "Малая Гостиная")
    assert "406" not in al                                                # пустой кабинет без подписи убран


def test_save_load_roundtrip_and_broken_file_falls_back_to_defaults(tmp_path):
    path = tmp_path / "a" / "room_aliases.json"
    room_aliases.save_aliases(str(path), {"305": {"label": "", "names": ["ф1"]}})
    assert room_aliases.load_aliases(str(path)) == {"305": {"label": "", "names": ["ф1"]}}
    path.write_text("{ не json", encoding="utf-8")
    assert room_aliases.load_aliases(str(path)) == room_aliases.DEFAULT_ALIASES


def _db_with_hall_lessons(tmp_path, monkeypatch, make_lesson):
    lessons = []
    for t in ["08:30", "09:20", "10:15"]:
        for k in range(8):
            lessons.append(make_lesson(day=1, start=t, teacher=f"П{k} П.П.", student=f"С{t}{k}", room=f"4{k}"))
    lessons.append(make_lesson(day=1, start="10:15", teacher="Иванов И.И.", student="Алексеев Ян", room="423"))
    lessons.append(make_lesson(day=1, start="10:15", teacher="Петров П.П.", student="Борисова Нина", room="малый зал"))
    lessons.append(make_lesson(day=1, start="10:15", teacher="Сидоров С.С.", student="Васильев Олег", room="м/ф"))   # малое фойе — отдельно
    conn = db.get_connection(str(tmp_path / "t.sqlite3"))
    db.save_import(conn, "imp", lessons, {})
    conn.commit()
    conn.close()
    monkeypatch.setattr(main_module, "DB_PATH", str(tmp_path / "t.sqlite3"))
    return TestClient(main_module.app)


def test_report_finds_hall_clash_and_old_imports_benefit(tmp_path, monkeypatch, make_lesson):
    client = _db_with_hall_lessons(tmp_path, monkeypatch, make_lesson)

    api = client.get("/api/report/imp").json()
    room_clashes = [c for c in api["conflicts"] if c["type"] == "room_double_booked"]
    assert len(room_clashes) == 1
    assert {room_clashes[0]["a"]["room"], room_clashes[0]["b"]["room"]} == {"423", "малый зал"}   # в файле — как написано


def test_layout_hall_and_foyer_are_different_columns(tmp_path, monkeypatch, make_lesson):
    client = _db_with_hall_lessons(tmp_path, monkeypatch, make_lesson)
    doc = Document(BytesIO(client.get("/report/imp/layout.docx", params={"rooms": ["423", "м/ф", "малый зал"]}).content))
    assert [c.text for c in doc.tables[0].rows[0].cells] == ["Часы", "Малый зал (423)", "Малое фойе (м/ф)"]


def test_layout_has_one_column_with_label_for_the_hall(tmp_path, monkeypatch, make_lesson):
    client = _db_with_hall_lessons(tmp_path, monkeypatch, make_lesson)

    doc = Document(BytesIO(client.get("/report/imp/layout.docx", params={"rooms": ["40", "423", "малый зал"]}).content))
    header = [c.text for c in doc.tables[0].rows[0].cells]
    assert header == ["Часы", "40", "Малый зал (423)"]                                      # «малый зал» влился в 423 — один столбец

    wb = openpyxl.load_workbook(BytesIO(client.get("/report/imp/layout.xlsx", params={"rooms": ["423"]}).content))
    assert wb["Вторник"].cell(2, 2).value == "Малый зал (423)"
    assert "⚠ НАКЛАДКА" in wb["Вторник"].cell(5, 2).value                                  # в 10:15 в зале двое


def test_cabinets_page_lists_and_edits_aliases(tmp_path, monkeypatch, make_lesson):
    client = _db_with_hall_lessons(tmp_path, monkeypatch, make_lesson)

    page = client.get("/report/imp/cabinets").text
    assert "Один кабинет — несколько названий" in page and "малый зал" in page and "Малое фойе" in page

    client.post("/report/imp/cabinets/alias/add", data={"name": "Фойе 2 эт.", "room": "305", "label": "Фойе"})
    data = json.loads(open(main_module.ALIASES_PATH, encoding="utf-8").read())["rooms"]
    assert data["305"] == {"label": "Фойе", "names": ["Фойе 2 эт."]} and "423" in data       # дефолт сохранён вместе с новым

    client.post("/report/imp/cabinets/alias/remove", data={"name": "малый зал"})
    assert "малый зал" not in json.loads(open(main_module.ALIASES_PATH, encoding="utf-8").read())["rooms"]["423"]["names"]


def test_registry_room_typed_as_alias_is_resolved_for_free_rooms(tmp_path, monkeypatch, make_lesson):
    from app import cabinets
    client = _db_with_hall_lessons(tmp_path, monkeypatch, make_lesson)
    config = cabinets.load_config(main_module.CABINETS_PATH)
    cabinets.add_room(config, "Малый зал")                                                  # в списке кафедры написали так
    cabinets.save_config(main_module.CABINETS_PATH, config)

    page = client.get("/report/imp/free-rooms?day=1").text

    assert ">423</a>" in page and "Малый зал" in page                                      # опознан как кабинет 423


def test_backup_keeps_aliases(tmp_path):
    cab, dis, al_path = str(tmp_path / "c.json"), str(tmp_path / "d.json"), str(tmp_path / "a.json")
    room_aliases.save_aliases(al_path, {"305": {"label": "Фойе", "names": ["ф2"]}})
    data = backup.build_backup(cab, dis, "0.3.8", al_path)
    assert data["room_aliases"]["305"]["names"] == ["ф2"]

    new_al = str(tmp_path / "new" / "a.json")
    backup.restore_backup(json.loads(json.dumps(data)), str(tmp_path / "n1.json"), str(tmp_path / "n2.json"), new_al)
    restored = room_aliases.load_aliases(new_al)
    assert restored["305"] == {"label": "Фойе", "names": ["ф2"]}
    assert restored["423"]["names"]                                                         # значения по умолчанию остались
