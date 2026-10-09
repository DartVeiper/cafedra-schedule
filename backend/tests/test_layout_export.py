"""Раскладка расписания по кабинетам — Word и Excel (как ручная «Раскладка …»)."""
from __future__ import annotations

from io import BytesIO

import openpyxl
from docx import Document
from fastapi.testclient import TestClient

from app import cabinets, db, layout_export
from app import main as main_module
from app.models import LessonType


def _week(make_lesson):
    """Вторник: сетка из 3 слотов (по 8 занятий в каждом — сетка считается надёжной)."""
    out = []
    for t in ["08:30", "09:20", "10:15"]:
        for k in range(8):
            out.append(make_lesson(day=1, start=t, teacher=f"Преп{k} П.П.", student=f"Студент{t}{k}", room=f"4{k}", subject="спец. инструмент"))
    return out


def test_teacher_and_accompanist_are_one_entry_with_concertmaster_in_parentheses(make_lesson):
    lessons = _week(make_lesson)
    lessons.append(make_lesson(day=1, start="08:30", teacher=None, accompanist="Шейна О.С.", student="Студент08:300", room="40",
                               subject="спец. инструмент"))

    layout = layout_export.build_layout(lessons, ["40"])

    cell = layout[0]["rows"][0]["cells"][0]
    assert cell["state"] == "busy" and len(cell["entries"]) == 1
    lines = layout_export.entry_lines(cell["entries"][0])
    assert lines[0] == "Преп0 П.П."                      # преподаватель — первая (жирная) строка
    assert lines[1].startswith("Студент08:300")
    assert lines[-1] == "Шейна О.С. (концертмейстер)"


def test_clash_group_lessons_and_curator_hour_in_layout(make_lesson):
    lessons = _week(make_lesson)
    lessons.append(make_lesson(day=1, start="08:30", teacher="Другой Д.Д.", student="Иной Студент", room="40"))  # накладка
    lessons.append(make_lesson(day=1, start="09:20", duration=90, lesson_type=LessonType.GROUP, student=None,
                               teacher="доц. Лектор Л.Л.", group="92Ф", subject="История", room="41"))
    special = [{"day": 1, "start": "14:25", "subject": "КУРАТОРСКИЙ ЧАС"}]

    layout = layout_export.build_layout(lessons, ["40", "41"], special)
    rows = layout[0]["rows"]

    assert rows[0]["cells"][0]["state"] == "clash"
    texts = layout_export._cell_text(rows[0]["cells"][0])
    assert texts.startswith("⚠ НАКЛАДКА") and "— — —" in texts
    group_first = layout_export.entry_lines(rows[1]["cells"][1]["entries"][0])
    assert group_first[0] == "ПАРА 92Ф" and group_first[2] == "Лектор Л.Л."          # звание убрано
    assert layout_export.entry_lines(rows[2]["cells"][1]["entries"][0]) == ["↑ пара продолжается"]
    assert rows[-1]["start"] == "14:25" and rows[-1]["special"] == "Кураторский час"
    assert rows[-1]["cells"][0]["state"] == "special"


def test_word_file_has_a_table_per_day_with_rooms_as_columns(make_lesson):
    lessons = _week(make_lesson) + [make_lesson(day=3, start=t, teacher=f"Ч{k} Ч.Ч.", student=f"С{t}{k}", room=f"4{k}")
                                    for t in ["08:30", "09:20"] for k in range(12)]
    data = layout_export.build_docx(lessons, ["40", "41", "42"], [], "Раскладка по кабинетам")

    doc = Document(BytesIO(data))
    assert len(doc.tables) == 2                                   # вторник и четверг
    header = [c.text for c in doc.tables[0].rows[0].cells]
    assert header == ["Часы", "40", "41", "42"]
    assert doc.tables[0].rows[1].cells[0].text == "08:30"
    assert "Преп0 П.П." in doc.tables[0].rows[1].cells[1].text
    assert any("Раскладка по кабинетам — Вторник" in p.text for p in doc.paragraphs)
    section = doc.sections[0]
    assert section.page_width > section.page_height               # альбомная страница


def test_excel_has_sheet_per_day_and_flat_list(make_lesson):
    lessons = _week(make_lesson)
    wb = openpyxl.load_workbook(BytesIO(layout_export.build_xlsx(lessons, ["40", "41"], [], "Раскладка")))

    assert wb.sheetnames == ["Вторник", "Все занятия"]
    day = wb["Вторник"]
    assert [day.cell(2, c).value for c in (1, 2, 3)] == ["Время", "40", "41"]
    assert "Преп0 П.П." in day.cell(3, 2).value
    flat = wb["Все занятия"]
    assert flat.max_row == 1 + len(lessons)                       # заголовок + все занятия (в т.ч. вне выбранных кабинетов)
    assert [c.value for c in flat[1]][:5] == ["День", "Начало", "Конец", "Кабинет", "Тип"]
    assert flat.auto_filter.ref is not None


def test_layout_rooms_use_registry_else_all_individual_rooms(make_lesson):
    lessons = _week(make_lesson)
    assert cabinets.layout_rooms(lessons, ["41", "40"]) == ["40", "41"]
    assert cabinets.layout_rooms(lessons, []) == [f"4{k}" for k in range(8)]


def test_layout_routes(tmp_path, monkeypatch, make_lesson):
    db_path = tmp_path / "t.sqlite3"
    conn = db.get_connection(str(db_path))
    db.save_import(conn, "imp", _week(make_lesson), {"special_slots": []})
    conn.commit()
    conn.close()
    monkeypatch.setattr(main_module, "DB_PATH", str(db_path))
    client = TestClient(main_module.app)

    word = client.get("/report/imp/layout.docx")
    assert word.status_code == 200 and word.content[:2] == b"PK"
    assert len(Document(BytesIO(word.content)).tables) == 1
    xl = client.get("/report/imp/layout.xlsx")
    assert xl.status_code == 200 and openpyxl.load_workbook(BytesIO(xl.content)).sheetnames == ["Вторник", "Все занятия"]
    assert client.get("/report/nope/layout.docx").status_code == 404
    assert client.get("/report/nope/layout.xlsx").status_code == 404
    assert "layout.docx" in client.get("/report/imp").text   # кнопка на странице отчёта
