"""Листы преподавателям: «ваши накладки» от лица каждого участника."""
from __future__ import annotations

from fastapi.testclient import TestClient

from app import db, teacher_sheets, timegrid
from app import main as main_module
from app.conflicts import find_conflicts
from app.models import LessonType


def _room_clash(make_lesson):
    return [
        make_lesson(day=1, start="14:35", teacher="Иванов И.И.", student="Алексеев Ян", room="421", subject="спец. инструмент"),
        make_lesson(day=1, start="14:35", teacher="доц. Петров П.П.", student="Борисова Нина", room="421", subject="ансамбль"),
    ]


def test_each_party_gets_the_conflict_in_first_person(make_lesson):
    sheets = teacher_sheets.build_sheets(find_conflicts(_room_clash(make_lesson)))

    assert set(sheets) == {"иванови.и.", "петровп.п."}
    ivanov = sheets["иванови.и."].items[0]
    assert ivanov.kind == "Кабинет занят дважды"
    assert "у вас — студент Алексеев Ян" in ivanov.text and "у Петров П.П." in ivanov.text  # звание отброшено
    assert "421" in ivanov.text
    petrov = sheets["петровп.п."].items[0]
    assert "у вас — студент Борисова Нина" in petrov.text and "у Иванов И.И." in petrov.text
    assert ivanov.when == "14:35–15:20" and ivanov.day == 1


def test_person_key_ignores_title_case_and_spaces():
    assert teacher_sheets.person_key("доц. Топорова Е. П.") == teacher_sheets.person_key("Топорова Е.П.")
    assert teacher_sheets.person_key(None) is None


def test_student_vs_group_text_for_teacher_and_group_teacher(make_lesson):
    ind = make_lesson(day=0, start="12:00", teacher="Ананьев А.А.", student="Шведов Никифор", group="92Ф", room="418")
    grp = make_lesson(day=0, start="12:00", teacher="Круговая Н.И.", student=None, group="92Ф", lesson_type=LessonType.GROUP,
                      duration=90, subject="Литература", room="корпус №5 ауд.316")
    conflicts = [c for c in find_conflicts([ind, grp]) if c.type.value == "student_vs_group"]

    sheets = teacher_sheets.build_sheets(conflicts)

    ananyev = sheets["ананьева.а."].items[0].text
    assert "Шведов Никифор в это время на групповом занятии (Литература, Круговая Н.И." in ananyev
    krugovaya = sheets["круговаян.и."].items[0].text
    assert "У вашей группы 92Ф" in krugovaya and "Шведов Никифор" in krugovaya


def test_review_flag_and_irregular_time_items(make_lesson):
    lessons = []
    for t in ["08:30", "09:20", "10:15"]:
        for k in range(8):
            lessons.append(make_lesson(day=1, start=t, teacher=f"П{k}", student=f"С{t}{k}", room=f"3{k}"))
    odd = make_lesson(day=1, start="10:22", teacher="Толмачева Э.Г.", student="Габриэлян София", room="417")
    lessons.append(odd)

    sheets = teacher_sheets.build_sheets([], timegrid.find_irregular_times(lessons))

    item = sheets["толмачеваэ.г."].items[0]
    assert item.kind == "Нестандартное время" and not item.certain and "10:22" in item.text


def test_all_people_lists_teachers_and_accompanists_only(make_lesson):
    people = teacher_sheets.all_people([
        make_lesson(teacher="Иванов И.И.", student="Студент"),
        make_lesson(teacher=None, accompanist="Смирнова А.А.", student="Студент"),
    ])
    assert sorted(people.values()) == ["Иванов И.И.", "Смирнова А.А."]


def _client(tmp_path, monkeypatch, make_lesson):
    db_path = tmp_path / "t.sqlite3"
    conn = db.get_connection(str(db_path))
    db.save_import(conn, "imp", _room_clash(make_lesson) + [make_lesson(day=2, teacher="Чистый Ч.Ч.", student="Ничей", room="999")], {})
    conn.commit()
    conn.close()
    monkeypatch.setattr(main_module, "DB_PATH", str(db_path))
    return TestClient(main_module.app)


def test_teacher_pages_and_word_export(tmp_path, monkeypatch, make_lesson):
    client = _client(tmp_path, monkeypatch, make_lesson)

    lst = client.get("/report/imp/teachers").text
    assert "Иванов И.И." in lst and "Петров П.П." in lst and "Без накладок (1)" in lst

    sheet = client.get("/report/imp/teacher", params={"name": "Иванов И.И."}).text
    assert "Накладки в расписании: Иванов И.И." in sheet and "у вас — студент Алексеев Ян" in sheet

    clean = client.get("/report/imp/teacher", params={"name": "Чистый Ч.Ч."}).text
    assert "Накладок не найдено" in clean

    everything = client.get("/report/imp/teachers/print").text
    assert everything.count('class="sheet"') == 2  # листы только у тех, у кого есть что показать

    assert client.get("/report/imp/teacher", params={"name": "Неизвестный"}).status_code == 404
    assert client.get("/report/nope/teachers").status_code == 404

    docx = client.get("/report/imp/teacher.docx", params={"name": "Иванов И.И."})
    assert docx.status_code == 200 and docx.content[:2] == b"PK"  # настоящий .docx (zip)
    from io import BytesIO
    from docx import Document
    d = Document(BytesIO(docx.content))
    assert "Иванов И.И." in d.paragraphs[0].text and len(d.tables) == 1


def test_dismissed_conflicts_do_not_appear_on_sheets(tmp_path, monkeypatch, make_lesson):
    client = _client(tmp_path, monkeypatch, make_lesson)
    from app.conflicts import conflict_key
    key = conflict_key(find_conflicts(_room_clash(make_lesson))[0])
    client.post("/report/imp/dismiss", data={"key": key})

    lst = client.get("/report/imp/teachers").text
    assert "Накладок не найдено ни у кого" in lst
