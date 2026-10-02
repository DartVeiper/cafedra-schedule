"""Парсер Формата 1: индивидуальные и мелкогрупповые занятия (Excel, .xls/.xlsx).

Один файл — один преподаватель ИЛИ концертмейстер (кто именно — определяется
по строке шапки листа "Преподаватель ... Концертмейстер ...", а не по имени файла:
на реальных данных кафедры встречаются файлы, названные в честь человека, который
в шапке листа указан как концертмейстер, а не как преподаватель, и наоборот).

Лист может быть пустым (0x0 у xlrd, если лист вообще не использовался), либо
содержать только шапку и пустую сетку (реальный бланк без данных), либо содержать
данные. Один файл может иметь данные сразу на нескольких листах книги —
это не дубликаты по умолчанию, но на практике они часто почти дублируют друг
друга (один лист — черновик, другой — актуальная версия); окончательное решение,
что с этим делать, принимает dedup.py на уровне всего пакета импорта.
"""
from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from datetime import time

import openpyxl
import xlrd

from app.models import Lesson, LessonType, ParseWarning, PersonRole, SourceRef, INDIVIDUAL_LESSON_MINUTES
from app.parsers.common import (
    DAY_NAME_TO_INDEX,
    excel_time_fraction_to_time,
    normalize_group,
    normalize_person_name,
    normalize_room,
    normalize_room_display,
    normalize_subject,
)

_HEADER_KAFEDRA_RE = re.compile(r"кафедра", re.IGNORECASE)
_HEADER_PREP_RE = re.compile(r"Преподаватель", re.IGNORECASE)
_HEADER_CONC_RE = re.compile(r"Концертмейстер", re.IGNORECASE)
_TRAILER_RE = re.compile(r"дата составления", re.IGNORECASE)
_FILENAME_TAIL_RE = re.compile(r"\s*\d+\s*семестр.*$", re.IGNORECASE)

COLUMNS_PER_DAY_BLOCK = 5  # Время | ФИО обучающегося | Группа | Предмет | Ауд.


@dataclass
class SheetParseResult:
    file_name: str
    sheet_name: str
    role: PersonRole | None
    person_name: str | None
    lessons: list[Lesson] = field(default_factory=list)
    warnings: list[ParseWarning] = field(default_factory=list)
    header_text: str | None = None

    @property
    def is_empty(self) -> bool:
        return not self.lessons and self.person_name is None


def parse_individual_workbook(path: str, known_groups: set[str] | None = None) -> list[SheetParseResult]:
    """Разбирает один файл Формата 1, возвращая результат по каждому листу книги."""
    file_name = os.path.basename(path)
    sheets = _load_sheets_as_grid(path)
    return [
        _parse_sheet(file_name, sheet_name, rows, known_groups)
        for sheet_name, rows in sheets.items()
    ]


def _load_sheets_as_grid(path: str) -> dict[str, list[list[object]]]:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xls":
        return _load_xls(path)
    if ext == ".xlsx":
        return _load_xlsx(path)
    raise ValueError(f"Неподдерживаемое расширение файла для Формата 1: {path}")


def _load_xls(path: str) -> dict[str, list[list[object]]]:
    book = xlrd.open_workbook(path)
    out: dict[str, list[list[object]]] = {}
    for name in book.sheet_names():
        sh = book.sheet_by_name(name)
        out[name] = [[sh.cell(r, c).value for c in range(sh.ncols)] for r in range(sh.nrows)]
    return out


def _load_xlsx(path: str) -> dict[str, list[list[object]]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    out: dict[str, list[list[object]]] = {}
    for ws in wb.worksheets:
        out[ws.title] = [[cell.value for cell in row] for row in ws.iter_rows()]
    return out


def _parse_sheet(
    file_name: str,
    sheet_name: str,
    rows: list[list[object]],
    known_groups: set[str] | None,
) -> SheetParseResult:
    warnings: list[ParseWarning] = []
    if not rows:
        return SheetParseResult(file_name, sheet_name, None, None, [], warnings, None)

    header_text = _find_header_text(rows)
    teacher, accompanist = (None, None)
    if header_text:
        teacher, accompanist = _extract_header_roles(header_text)

    role: PersonRole | None = None
    person_name: str | None = None
    if teacher:
        role, person_name = PersonRole.TEACHER, teacher
    elif accompanist:
        role, person_name = PersonRole.ACCOMPANIST, accompanist

    if header_text and person_name is None:
        warnings.append(
            ParseWarning(
                file_name,
                sheet_name,
                "В шапке листа не найдено ФИО ни в поле 'Преподаватель', ни в поле "
                "'Концертмейстер' — лист пропущен",
            )
        )
    elif person_name is not None:
        _warn_if_filename_mismatch(file_name, sheet_name, person_name, warnings)

    lessons: list[Lesson] = []
    if person_name is not None:
        blocks = _find_day_header_rows(rows)
        for i, (row_idx, day_cols) in enumerate(blocks):
            data_start = row_idx + 2  # + строка дней, + строка заголовков колонок
            data_end = blocks[i + 1][0] if i + 1 < len(blocks) else len(rows)
            for r in range(data_start, min(data_end, len(rows))):
                row = rows[r]
                if _is_trailer_row(row):
                    break
                for col_start, day_idx in day_cols.items():
                    lesson = _extract_lesson_from_row(
                        row, col_start, day_idx, file_name, sheet_name, r,
                        role, person_name, known_groups, warnings,
                    )
                    if lesson:
                        lessons.append(lesson)

    return SheetParseResult(file_name, sheet_name, role, person_name, lessons, warnings, header_text)


def _guess_name_from_filename(file_name: str) -> str | None:
    stem = os.path.splitext(file_name)[0]
    stem = _FILENAME_TAIL_RE.sub("", stem).strip()
    return stem or None


def _warn_if_filename_mismatch(
    file_name: str, sheet_name: str, person_name: str, warnings: list[ParseWarning]
) -> None:
    """Имя файла — не источник истины (см. модуль), но расхождение с шапкой часто
    выдаёт опечатку в исходных данных (наблюдалось: файл 'Черняева В.А....xlsx',
    а в шапке листа 'Концертмейстер Черняева В.О.') — стоит показать методисту."""
    guess = _guess_name_from_filename(file_name)
    if not guess:
        return
    # Сравниваем без пробелов, чтобы не шуметь на форматирование ('Е.П.' vs 'Е. П.'),
    # но ловить реальные расхождения в фамилии/инициалах.
    guess_key = re.sub(r"\s+", "", guess).lower()
    person_key = re.sub(r"\s+", "", person_name).lower()
    if guess_key and person_key and guess_key != person_key:
        warnings.append(
            ParseWarning(
                file_name, sheet_name,
                f"Имя файла ('{guess}') не совпадает с ФИО в шапке листа "
                f"('{person_name}') — возможно, опечатка в одном из них",
            )
        )


def _find_header_text(rows: list[list[object]]) -> str | None:
    for row in rows[:6]:
        for val in row:
            if isinstance(val, str) and _HEADER_KAFEDRA_RE.search(val):
                return val
    return None


def _extract_header_roles(header_text: str) -> tuple[str | None, str | None]:
    prep_match = _HEADER_PREP_RE.search(header_text)
    if not prep_match:
        return None, None
    conc_match = _HEADER_CONC_RE.search(header_text)
    if conc_match:
        teacher_part = header_text[prep_match.end() : conc_match.start()]
        accompanist_part = header_text[conc_match.end() :]
    else:
        teacher_part = header_text[prep_match.end() :]
        accompanist_part = ""
    return normalize_person_name(teacher_part), normalize_person_name(accompanist_part)


def _find_day_header_rows(rows: list[list[object]]) -> list[tuple[int, dict[int, int]]]:
    """Строки с названиями дней недели -> {колонка_начала_блока: индекс_дня}."""
    blocks: list[tuple[int, dict[int, int]]] = []
    for r, row in enumerate(rows):
        day_cols: dict[int, int] = {}
        for c, val in enumerate(row):
            if isinstance(val, str):
                idx = DAY_NAME_TO_INDEX.get(val.strip().lower())
                if idx is not None:
                    day_cols[c] = idx
        if day_cols:
            blocks.append((r, day_cols))
    return blocks


def _is_trailer_row(row: list[object]) -> bool:
    for val in row:
        if isinstance(val, str) and _TRAILER_RE.search(val):
            return True
    return False


def _normalize_time_cell(value: object) -> time | None:
    if isinstance(value, time):
        return value
    if isinstance(value, (int, float)) and 0 <= value < 1:
        return excel_time_fraction_to_time(value)
    return None


def _extract_lesson_from_row(
    row: list[object],
    col_start: int,
    day_idx: int,
    file_name: str,
    sheet_name: str,
    row_idx: int,
    role: PersonRole | None,
    person_name: str | None,
    known_groups: set[str] | None,
    warnings: list[ParseWarning],
) -> Lesson | None:
    if col_start + COLUMNS_PER_DAY_BLOCK > len(row):
        return None
    time_val, student_val, group_val, subject_val, room_val = row[col_start : col_start + COLUMNS_PER_DAY_BLOCK]

    student = normalize_person_name(student_val) if isinstance(student_val, str) else None
    if not student:
        return None  # незанятый тайм-слот

    lesson_time = _normalize_time_cell(time_val)
    if lesson_time is None:
        warnings.append(
            ParseWarning(
                file_name, sheet_name,
                f"Не удалось прочитать время занятия у студента '{student}'", row_idx,
            )
        )
        return None

    group_raw = None if group_val in (None, "") else str(group_val).strip()
    group_norm = normalize_group(group_raw)
    if group_norm and known_groups is not None and group_norm not in known_groups:
        warnings.append(
            ParseWarning(
                file_name, sheet_name,
                f"Незнакомый код группы '{group_raw}' у студента '{student}' — "
                "возможно, опечатка (сверьте со списком групп)", row_idx,
            )
        )

    subject = normalize_subject(subject_val) if isinstance(subject_val, str) else None
    room_raw = normalize_room_display(room_val)
    room_norm = normalize_room(room_val)

    return Lesson(
        lesson_type=LessonType.INDIVIDUAL,
        day_of_week=day_idx,
        start_time=lesson_time,
        duration_minutes=INDIVIDUAL_LESSON_MINUTES,
        teacher_name=person_name if role == PersonRole.TEACHER else None,
        accompanist_name=person_name if role == PersonRole.ACCOMPANIST else None,
        student_name=student,
        group_raw=group_raw,
        group_normalized=group_norm,
        subject=subject,
        room_raw=room_raw,
        room_normalized=room_norm,
        source=SourceRef(file_name=file_name, sheet_name=sheet_name, row_index=row_idx),
    )
