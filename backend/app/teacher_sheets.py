"""Листы для преподавателей: «ваши накладки» одним печатным листом на человека.

Методист не показывает преподавателю весь отчёт — он пишет каждому, что именно
у него не так. Здесь для каждого преподавателя/концертмейстера собираются все
накладки, где он участвует (своим занятием или как «другая сторона»), и пишутся
понятными фразами от его лица: «у вас — …; у Петрова — …».
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field

from app import timegrid
from app.models import DAY_NAMES_RU, Conflict, ConflictType, Lesson, LessonType
from app.parsers.common import strip_academic_title

_UNDER_KIND = {
    ConflictType.ROOM_DOUBLE_BOOKED: "Кабинет занят дважды",
    ConflictType.STUDENT_DOUBLE_BOOKED: "Студент в двух местах",
    ConflictType.STUDENT_VS_GROUP: "Студент: индивидуальное и групповое",
    ConflictType.TEACHER_DOUBLE_BOOKED: "Вы в двух местах",
}


@dataclass
class SheetItem:
    day: int
    start: int                 # минуты — для сортировки
    when: str                  # "14:35–15:20"
    kind: str
    text: str
    certain: bool = True
    note: str | None = None


@dataclass
class Sheet:
    key: str
    name: str
    items: list[SheetItem] = field(default_factory=list)

    @property
    def certain_count(self) -> int:
        return sum(1 for i in self.items if i.certain)

    @property
    def review_count(self) -> int:
        return sum(1 for i in self.items if not i.certain)


def person_key(name: str | None) -> str | None:
    """Ключ человека: без звания, регистра и пробелов ('Топорова Е. П.' == 'Топорова Е.П.')."""
    if not name:
        return None
    return re.sub(r"\s+", "", strip_academic_title(name)).casefold()


def _lesson_person(l: Lesson) -> tuple[str | None, str | None]:
    name = l.teacher_name or l.accompanist_name
    return person_key(name), (strip_academic_title(name) if name else None)


def _who(l: Lesson) -> str:
    return _lesson_person(l)[1] or "—"


def _room(l: Lesson) -> str:
    r = re.sub(r"\s+", " ", (l.room_raw or "")).strip()
    if not r:
        return ""
    low = r.lower()
    return r if ("ауд" in low or "корпус" in low or "с/з" in low) else f"ауд. {r}"


def _fmt(l: Lesson) -> str:
    return timegrid.fmt_minutes(l.start_minutes) + "–" + timegrid.fmt_minutes(l.end_minutes)


def _desc(l: Lesson) -> str:
    bits = []
    if l.student_name:
        bits.append(f"студент {l.student_name}")
    if l.subject:
        bits.append(l.subject)
    if l.group_raw and not l.student_name:
        bits.append(f"гр. {l.group_raw}")
    room = _room(l)
    if room:
        bits.append(room)
    return ", ".join(bits) or "занятие"


def _overlap(a: Lesson, b: Lesson) -> tuple[int, int]:
    return max(a.start_minutes, b.start_minutes), min(a.end_minutes, b.end_minutes)


def _text(c: Conflict, mine: Lesson, other: Lesson) -> str:
    who_o = _who(other)
    if c.type == ConflictType.ROOM_DOUBLE_BOOKED:
        room = _room(mine) or _room(other) or "кабинет"
        room = ("Кабинет " + room[5:]) if room.startswith("ауд. ") else (room[0].upper() + room[1:])
        return f"{room} занят одновременно: у вас — {_desc(mine)}; у {who_o} — {_desc(other)}."
    if c.type == ConflictType.STUDENT_DOUBLE_BOOKED:
        student = mine.student_name or other.student_name
        return f"Студент {student} записан в двух местах сразу: у вас — {_desc(mine)}; у {who_o} — {_desc(other)}."
    if c.type == ConflictType.STUDENT_VS_GROUP:
        if mine.lesson_type == LessonType.INDIVIDUAL:
            return (f"Студент {mine.student_name} в это время на групповом занятии "
                    f"({other.subject or 'пара'}, {who_o}, {_room(other) or 'аудитория не указана'}); у вас — {_desc(mine)}.")
        return (f"У вашей группы {mine.group_raw or ''} в это время пара ({mine.subject or 'занятие'}), а у студента "
                f"{other.student_name} индивидуальное занятие у {who_o} ({_desc(other)}).")
    return f"Вы записаны одновременно в двух местах: {_desc(mine)} и {_desc(other)}."


def build_sheets(
    conflicts: list[Conflict], irregular: list | None = None, all_lessons: list[Lesson] | None = None
) -> dict[str, Sheet]:
    """{ключ человека: Sheet}. conflicts — только активные (не помеченные «не накладка»)
    и без пар «преп.+концертмейстер»; irregular — timegrid.find_irregular_times()."""
    sheets: dict[str, Sheet] = {}

    def sheet(key: str, name: str) -> Sheet:
        return sheets.setdefault(key, Sheet(key, name))

    for c in conflicts:
        if c.type == ConflictType.ACCOMPANIST_PAIRING:
            continue
        a, b = c.lesson_a, c.lesson_b
        ka, na = _lesson_person(a)
        kb, nb = _lesson_person(b)
        start, end = _overlap(a, b)
        when = f"{timegrid.fmt_minutes(start)}–{timegrid.fmt_minutes(end)}"
        parties = []
        if ka:
            parties.append((ka, na, a, b))
        if kb and kb != ka:
            parties.append((kb, nb, b, a))
        for key, name, mine, other in parties:
            sheet(key, name).items.append(SheetItem(
                c.day_of_week, start, when, _UNDER_KIND[c.type], _text(c, mine, other), c.is_certain, c.note,
            ))

    for item in irregular or []:
        key, name = _lesson_person(item.lesson)
        if key:
            l = item.lesson
            sheet(key, name).items.append(SheetItem(
                l.day_of_week, l.start_minutes, _fmt(l), "Нестандартное время",
                f"{_desc(l)}: {timegrid.describe(item)}.", certain=False,
            ))

    for s in sheets.values():
        s.items.sort(key=lambda i: (i.day, i.start, i.kind))
    return sheets


def all_people(lessons: list[Lesson]) -> dict[str, str]:
    """{ключ: имя} всех преподавателей и концертмейстеров — чтобы показать и тех, у кого накладок нет."""
    out: dict[str, str] = {}
    for l in lessons:
        key, name = _lesson_person(l)
        if key and key not in out:
            out[key] = name
    return out


def build_docx(sheet: Sheet, created_text: str):
    """Лист одного преподавателя в Word — чтобы отправить файлом."""
    from docx import Document

    doc = Document()
    doc.add_heading(f"Накладки в расписании: {sheet.name}", level=1)
    doc.add_paragraph(f"Составлено {created_text}. Найдено: {sheet.certain_count} явных, "
                      f"{sheet.review_count} требуют вашей проверки.")
    if not sheet.items:
        doc.add_paragraph("Накладок не найдено — всё в порядке.")
        return doc
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "День и время", "Что не так", "Примечание"
    for it in sheet.items:
        row = table.add_row().cells
        row[0].text = f"{DAY_NAMES_RU[it.day]}, {it.when}"
        row[1].text = f"{it.kind}. {it.text}"
        row[2].text = ("[проверить] " if not it.certain else "") + (it.note or "")
    return doc
