"""FastAPI-приложение: загрузка файлов расписания, поиск накладок, отчёт в браузере.

Доступно в локальной сети кафедры без установки на компьютеры пользователей —
достаточно открыть http://<адрес-сервера>:8000 в браузере.
"""
from __future__ import annotations

import io
import json
import os
import re
import shutil
import uuid
from datetime import datetime
from urllib.parse import quote, urlencode

from fastapi import FastAPI, File, Form, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates

from app import cabinets as cabinets_module
from app import db, self_update
from app import dismissed as dismissed_module
from app import layout_export, load_checks, teacher_sheets, timegrid
from app.conflicts import conflict_key, find_conflicts
from app.models import DAY_NAMES_RU, ConflictType
from app.parsers.common import strip_academic_title
from app.paths import app_package_dir, cabinets_config_path, dismissed_conflicts_path, migrate_legacy_storage, runtime_data_dir
from app.pipeline import known_groups_from_filenames, load_group_lessons, load_individual_lessons
from app import backup as backup_module
from app import changelog
from app import compare as compare_module
from app.mdlite import render_markdown_lite
from app.update_check import cached_update, check_for_update
from app.version import APP_VERSION
from app.zip_utils import extract_zip

APP_DIR = app_package_dir()
migrate_legacy_storage()  # переносит data/ и config/ со старого места (рядом с .exe), если ещё там
DATA_DIR = runtime_data_dir()
UPLOADS_DIR = os.path.join(DATA_DIR, "uploads")
DB_PATH = os.path.join(DATA_DIR, "db", "cafedra.sqlite3")
CABINETS_PATH = cabinets_config_path()
DISMISSED_PATH = dismissed_conflicts_path()

INDIVIDUAL_EXTS = {".xls", ".xlsx"}
GROUP_EXTS = {".doc", ".docx"}

_MONTHS_RU = [
    "", "января", "февраля", "марта", "апреля", "мая", "июня",
    "июля", "августа", "сентября", "октября", "ноября", "декабря",
]

os.makedirs(UPLOADS_DIR, exist_ok=True)
os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)

app = FastAPI(title="Проверка расписания — поиск накладок и сверка")
app.mount("/static", StaticFiles(directory=os.path.join(APP_DIR, "static")), name="static")
templates = Jinja2Templates(directory=os.path.join(APP_DIR, "templates"))


@app.exception_handler(Exception)
async def unhandled_exception_handler(request: Request, exc: Exception) -> HTMLResponse:
    """Страховка на случай непредвиденной ошибки — вместо голой 'Internal Server
    Error' показываем понятную страницу. Отдельные файлы при импорте и так не
    должны сюда попадать (см. pipeline.py), это именно последний рубеж."""
    return HTMLResponse(
        "<div style='font-family:sans-serif;max-width:640px;margin:60px auto;padding:0 20px'>"
        "<h1>Что-то пошло не так</h1>"
        "<p>Произошла непредвиденная ошибка. Попробуйте повторить действие; "
        "если это случилось при загрузке файлов — проверьте, что все файлы "
        "открываются в Excel/Word и не защищены паролем.</p>"
        f"<p style='color:#888;font-size:0.85em'>{type(exc).__name__}: {exc}</p>"
        "<p><a href='/'>На главную</a></p>"
        "</div>",
        status_code=500,
    )

CONFLICT_LABELS = {
    ConflictType.TEACHER_DOUBLE_BOOKED: "Преподаватель/концертмейстер в двух местах",
    ConflictType.ROOM_DOUBLE_BOOKED: "Аудитория занята дважды",
    ConflictType.STUDENT_DOUBLE_BOOKED: "Студент на двух индивидуальных одновременно",
    ConflictType.STUDENT_VS_GROUP: "Студент: индивидуальное пересекается с групповым",
    ConflictType.ACCOMPANIST_PAIRING: "Проверьте концертмейстера (не накладки)",
}


def _format_ru_datetime(value: str) -> str:
    """"2026-09-30 14:05:00" -> "30 сентября 2026, 14:05". Если формат
    неожиданный — возвращаем как есть, не роняя страницу из-за мелочи."""
    try:
        dt = datetime.strptime(value, "%Y-%m-%d %H:%M:%S")
    except (ValueError, TypeError):
        return value
    return f"{dt.day} {_MONTHS_RU[dt.month]} {dt.year}, {dt.strftime('%H:%M')}"


@app.get("/backup/download")
def backup_download():
    data = backup_module.build_backup(CABINETS_PATH, DISMISSED_PATH, APP_VERSION)
    body = json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8")
    name = f"cafedra-settings-{datetime.now():%Y-%m-%d}.json"
    return StreamingResponse(
        io.BytesIO(body), media_type="application/json",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )


@app.post("/backup/restore")
async def backup_restore(file: UploadFile = File(...)):
    raw = await file.read(backup_module.MAX_BACKUP_BYTES + 1)
    try:
        if len(raw) > backup_module.MAX_BACKUP_BYTES:
            raise backup_module.BackupError("Файл слишком большой для копии настроек.")
        try:
            data = json.loads(raw.decode("utf-8-sig"))
        except (ValueError, UnicodeDecodeError):
            raise backup_module.BackupError("Файл не читается — это не копия настроек программы.")
        result = backup_module.restore_backup(data, CABINETS_PATH, DISMISSED_PATH)
        msg = f"Копия восстановлена: добавлено кабинетов — {result['rooms']}, пометок «это не накладка» — {result['marks']}."
        qs = urlencode({"backup_msg": msg, "backup_ok": "1"})
    except backup_module.BackupError as e:
        qs = urlencode({"backup_msg": str(e), "backup_ok": ""})
    return RedirectResponse(url=f"/?{qs}#backup", status_code=303)


@app.get("/", response_class=HTMLResponse)
def upload_form(request: Request, backup_msg: str = "", backup_ok: str = ""):
    conn = db.get_connection(DB_PATH)
    try:
        recent = db.list_imports(conn)[:10]
    finally:
        conn.close()
    recent_display = [{"id": r["id"], "created_at": _format_ru_datetime(r["created_at"])} for r in recent]
    return templates.TemplateResponse(request, "upload.html", {
        "recent": recent_display,
        "update_info": check_for_update(),
        "can_self_update": self_update.can_self_update(),
        "app_version": APP_VERSION,
        "backup_msg": backup_msg[:300],
        "backup_ok": bool(backup_ok),
    })


@app.get("/api/version")
def api_version():
    """Версия запущенной программы — по ней страница обновления понимает, что новая версия
    запустилась после перезапуска (см. update_progress.html)."""
    return {"version": APP_VERSION}


@app.post("/update/apply", response_class=HTMLResponse)
def apply_update(request: Request):
    """Скачивает и устанавливает обновление поверх текущего .exe (см. self_update.py).
    Ссылку на скачивание берём из собственной кэшированной проверки, а не от
    клиента — чтобы нельзя было подсунуть форме произвольный URL для скачивания."""
    if not self_update.can_self_update():
        return HTMLResponse(
            "<p>Самообновление доступно только в собранной версии .exe на Windows. "
            "Для запуска из исходников используйте <code>git pull</code> "
            "(см. «Обновить и запустить.bat»).</p><p><a href='/'>Назад</a></p>",
            status_code=400,
        )
    info = check_for_update(force=True)
    if not info or not info.get("download_url"):
        return HTMLResponse(
            "<p>Обновление не найдено — возможно, его уже установили, или к "
            "релизу на GitHub не приложен .exe-файл.</p><p><a href='/'>Назад</a></p>",
            status_code=400,
        )
    try:
        self_update.start_update(info["download_url"], expected_size=info.get("size"))
    except Exception as e:
        return HTMLResponse(
            f"<p>Не удалось скачать обновление: {type(e).__name__}: {e}. "
            "Проверьте подключение к интернету и попробуйте ещё раз.</p>"
            "<p><a href='/'>Назад</a></p>",
            status_code=500,
        )
    self_update.schedule_exit()
    return templates.TemplateResponse(request, "update_progress.html", {
        "version": info["version"], "release_url": info["url"],
    })


@app.post("/import")
async def do_import(request: Request, files: list[UploadFile] = File(...)):
    import_id = uuid.uuid4().hex[:12]
    work_dir = os.path.join(UPLOADS_DIR, import_id)
    individual_dir = os.path.join(work_dir, "individual")
    group_dir = os.path.join(work_dir, "group")
    group_docx_dir = os.path.join(work_dir, "group_docx")
    zip_staging_dir = os.path.join(work_dir, "_zip_extracted")
    for d in (individual_dir, group_dir, group_docx_dir, zip_staging_dir):
        os.makedirs(d, exist_ok=True)

    upload_warnings: list[str] = []
    zip_rename_notes: list[str] = []

    for uf in files:
        # Только имя файла, без пути: имя приходит от браузера и попадает в os.path.join.
        name = os.path.basename((uf.filename or "").replace("\\", "/")) or "файл_без_имени"
        ext = os.path.splitext(name)[1].lower()
        raw_bytes = await uf.read()

        if ext == ".zip":
            tmp_zip_path = os.path.join(work_dir, name)
            with open(tmp_zip_path, "wb") as f:
                f.write(raw_bytes)
            try:
                result = extract_zip(tmp_zip_path, zip_staging_dir)
            except Exception as e:
                upload_warnings.append(
                    f"Архив '{name}' повреждён или это не ZIP-файл — пропущен целиком "
                    f"({type(e).__name__}: {e})"
                )
                continue
            zip_rename_notes.extend(result.renamed_notes)
            for extracted_path in result.extracted_paths:
                ext2 = os.path.splitext(extracted_path)[1].lower()
                if ext2 in INDIVIDUAL_EXTS:
                    shutil.copy(extracted_path, os.path.join(individual_dir, os.path.basename(extracted_path)))
                elif ext2 in GROUP_EXTS:
                    shutil.copy(extracted_path, os.path.join(group_dir, os.path.basename(extracted_path)))
                else:
                    upload_warnings.append(f"В архиве {name}: файл '{os.path.basename(extracted_path)}' — неизвестный формат, пропущен")
        elif ext in INDIVIDUAL_EXTS:
            with open(os.path.join(individual_dir, name), "wb") as f:
                f.write(raw_bytes)
        elif ext in GROUP_EXTS:
            with open(os.path.join(group_dir, name), "wb") as f:
                f.write(raw_bytes)
        else:
            upload_warnings.append(f"Файл '{name}' — неизвестный формат (ожидались .zip/.xls/.xlsx/.doc/.docx), пропущен")

    known_groups = known_groups_from_filenames(group_dir) if os.listdir(group_dir) else set()

    individual_report = None
    if os.listdir(individual_dir):
        individual_report = load_individual_lessons(individual_dir, known_groups=known_groups)

    group_lessons, group_warnings, group_failed_files = [], [], []
    special_slots: list[dict] = []
    if os.listdir(group_dir):
        group_lessons, group_warnings, group_failed_files = load_group_lessons(
            group_dir, group_docx_dir, special_slots_out=special_slots
        )

    all_lessons = list(individual_report.lessons if individual_report else []) + group_lessons

    failed_files = (individual_report.failed_files if individual_report else []) + group_failed_files

    notes = {
        "upload_warnings": upload_warnings,
        "zip_rename_notes": zip_rename_notes,
        "failed_files": failed_files,
        "parse_warnings": [w.label() for w in (individual_report.warnings if individual_report else [])] + [w.label() for w in group_warnings],
        "dropped_sheet_notes": individual_report.dropped_sheet_notes if individual_report else [],
        "merge_notes": individual_report.merge_notes if individual_report else [],
        "skipped_sheet_notes": individual_report.skipped_sheet_notes if individual_report else [],
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "special_slots": special_slots,  # кураторский час и т.п. — для подписей в сетке кабинета
        "individual_files": sorted(os.listdir(individual_dir)),
        "group_files": sorted(os.listdir(group_dir)),
    }

    conn = db.get_connection(DB_PATH)
    try:
        db.save_import(conn, import_id, all_lessons, notes)
        conn.commit()
    finally:
        conn.close()

    # Всё нужное (список занятий, имена файлов, предупреждения) уже в БД —
    # сырые загруженные файлы (реальные ФИО студентов/преподавателей) больше
    # не нужны и не должны бессрочно копиться на диске методиста.
    shutil.rmtree(work_dir, ignore_errors=True)

    return RedirectResponse(url=f"/report/{import_id}", status_code=303)


@app.post("/report/{import_id}/delete")
def delete_import(import_id: str):
    conn = db.get_connection(DB_PATH)
    try:
        db.delete_import(conn, import_id)
        conn.commit()
    finally:
        conn.close()
    return RedirectResponse(url="/", status_code=303)


def _load_lessons_or_none(import_id: str) -> list | None:
    conn = db.get_connection(DB_PATH)
    try:
        if not db.import_exists(conn, import_id):
            return None
        return db.load_lessons(conn, import_id)
    finally:
        conn.close()


@app.get("/report/{import_id}/cabinets", response_class=HTMLResponse)
def cabinets_page(request: Request, import_id: str):
    lessons = _load_lessons_or_none(import_id)
    if lessons is None:
        return HTMLResponse("<h1>Импорт не найден</h1><p><a href='/'>На главную</a></p>", status_code=404)

    config = cabinets_module.load_config(CABINETS_PATH)
    registered = cabinets_module.get_rooms(config)
    suggested = sorted(cabinets_module.suggest_rooms(lessons) - set(registered), key=cabinets_module.room_sort_key)

    return templates.TemplateResponse(request, "cabinets.html", {
        "import_id": import_id,
        "rooms": registered,
        "suggested_rooms": suggested,
        "suggested_floors": cabinets_module.group_rooms_by_floor(suggested),
        "irregular_by_room": cabinets_module.irregular_counts_by_room(lessons),
    })


@app.post("/report/{import_id}/cabinets/add")
def cabinets_add(import_id: str, room: str = Form(...)):
    config = cabinets_module.load_config(CABINETS_PATH)
    cabinets_module.add_room(config, room)
    cabinets_module.save_config(CABINETS_PATH, config)
    return RedirectResponse(url=f"/report/{import_id}/cabinets", status_code=303)


@app.post("/report/{import_id}/cabinets/add-many")
def cabinets_add_many(import_id: str, rooms: list[str] = Form(default=[])):
    config = cabinets_module.load_config(CABINETS_PATH)
    for room in rooms[:500]:
        cabinets_module.add_room(config, room)
    cabinets_module.save_config(CABINETS_PATH, config)
    return RedirectResponse(url=f"/report/{import_id}/cabinets", status_code=303)


@app.post("/report/{import_id}/cabinets/remove")
def cabinets_remove(import_id: str, room: str = Form(...)):
    config = cabinets_module.load_config(CABINETS_PATH)
    cabinets_module.remove_room(config, room)
    cabinets_module.save_config(CABINETS_PATH, config)
    return RedirectResponse(url=f"/report/{import_id}/cabinets", status_code=303)


@app.get("/report/{import_id}/cabinets/{room}", response_class=HTMLResponse)
def cabinet_schedule(request: Request, import_id: str, room: str):
    lessons, notes = _load_import(import_id)
    if lessons is None:
        return HTMLResponse("<h1>Импорт не найден</h1><p><a href='/'>На главную</a></p>", status_code=404)

    special_slots = (notes or {}).get("special_slots", [])
    schedule = cabinets_module.build_room_schedule(lessons, room, special_slots)
    return templates.TemplateResponse(request, "cabinet_schedule.html", {
        "import_id": import_id,
        "room": room,
        "grid_rows": cabinets_module.grid_rows(schedule, sorted(schedule.keys())),
        "break_rows": cabinets_module.break_rows(schedule, sorted(schedule.keys())),
        "day_notes": timegrid.day_notes(timegrid.build_day_grids(lessons), special_slots),
        "days_present": sorted(schedule.keys()),
        "irregular_here": [
            timegrid.describe(i) for i in timegrid.find_irregular_times(lessons)
            if i.lesson.room_normalized == room
        ],
        "day_names": DAY_NAMES_RU,
    })


def _today_text() -> str:
    d = datetime.now()
    return f"{d.day} {_MONTHS_RU[d.month]} {d.year}"


def _load_sheets(import_id: str):
    """(sheets, people) для листов преподавателям или None, если проверки нет."""
    lessons, _ = _load_import(import_id)
    if lessons is None:
        return None
    sel = _select_conflicts(lessons, _Filters())
    sheets = teacher_sheets.build_sheets(sel.conflicts, timegrid.find_irregular_times(lessons))
    return sheets, teacher_sheets.all_people(lessons)


_NOT_FOUND = HTMLResponse("<h1>Импорт не найден</h1><p><a href='/'>На главную</a></p>", status_code=404)


@app.get("/report/{import_id}/teachers", response_class=HTMLResponse)
def teachers_page(request: Request, import_id: str):
    loaded = _load_sheets(import_id)
    if loaded is None:
        return _NOT_FOUND
    sheets, people = loaded
    with_issues = sorted(sheets.values(), key=lambda s: (-len(s.items), s.name.casefold()))
    clean = sorted((n for k, n in people.items() if k not in sheets), key=str.casefold)
    return templates.TemplateResponse(request, "teachers.html", {
        "import_id": import_id, "with_issues": with_issues, "clean": clean,
    })


def _sheet_response(request: Request, import_id: str, names: list | None):
    loaded = _load_sheets(import_id)
    if loaded is None:
        return _NOT_FOUND
    sheets, people = loaded
    if names is None:                                   # все листы, у кого есть что показать
        chosen = sorted(sheets.values(), key=lambda s: s.name.casefold())
    else:
        chosen = []
        for name in names:
            key = teacher_sheets.person_key(name)
            if key in sheets:
                chosen.append(sheets[key])
            elif key in people:
                chosen.append(teacher_sheets.Sheet(key, people[key]))
        if not chosen:
            return HTMLResponse("<h1>Преподаватель не найден</h1><p><a href='javascript:history.back()'>Назад</a></p>", status_code=404)
    return templates.TemplateResponse(request, "teacher_sheet.html", {
        "import_id": import_id, "sheets": chosen, "single": names is not None,
        "created": _today_text(), "day_names": DAY_NAMES_RU,
    })


@app.get("/report/{import_id}/teacher", response_class=HTMLResponse)
def teacher_sheet_page(request: Request, import_id: str, name: str = ""):
    return _sheet_response(request, import_id, [name])


@app.get("/report/{import_id}/teachers/print", response_class=HTMLResponse)
def teachers_print_page(request: Request, import_id: str):
    return _sheet_response(request, import_id, None)


@app.get("/report/{import_id}/teacher.docx")
def teacher_sheet_docx(import_id: str, name: str = ""):
    loaded = _load_sheets(import_id)
    if loaded is None:
        return _NOT_FOUND
    sheets, people = loaded
    key = teacher_sheets.person_key(name)
    if key not in people:
        return HTMLResponse("Преподаватель не найден", status_code=404)
    doc = teacher_sheets.build_docx(sheets.get(key) or teacher_sheets.Sheet(key, people[key]), _today_text())
    import io
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename*=UTF-8''" + quote(f"nakladki_{name}.docx")},
    )


def _layout_inputs(import_id: str):
    lessons, notes = _load_import(import_id)
    if lessons is None:
        return None
    registry = cabinets_module.get_rooms(cabinets_module.load_config(CABINETS_PATH))
    return lessons, cabinets_module.layout_rooms(lessons, registry), (notes or {}).get("special_slots", [])


@app.get("/report/{import_id}/layout.docx")
def layout_docx(import_id: str):
    inputs = _layout_inputs(import_id)
    if inputs is None:
        return _NOT_FOUND
    lessons, rooms, special = inputs
    data = layout_export.build_docx(lessons, rooms, special, "Раскладка по кабинетам")
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename=raskladka_{import_id}.docx"},
    )


@app.get("/report/{import_id}/layout.xlsx")
def layout_xlsx(import_id: str):
    inputs = _layout_inputs(import_id)
    if inputs is None:
        return _NOT_FOUND
    lessons, rooms, special = inputs
    data = layout_export.build_xlsx(lessons, rooms, special, "Раскладка по кабинетам")
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=raskladka_{import_id}.xlsx"},
    )


@app.get("/report/{import_id}/load", response_class=HTMLResponse)
def load_page(request: Request, import_id: str, max_lessons: int = load_checks.MAX_LESSONS_PER_DAY,
              window_hours: float = load_checks.BIG_WINDOW_MINUTES / 60, day: str = ""):
    lessons, _ = _load_import(import_id)
    if lessons is None:
        return _NOT_FOUND
    max_lessons = min(max(max_lessons, 2), 20)
    window_hours = min(max(window_hours, 1), 12)
    items = load_checks.find_load_issues(lessons, max_lessons, int(window_hours * 60))
    chosen = _parse_day(day)
    shown = [i for i in items if chosen is None or i.day == chosen]
    return templates.TemplateResponse(request, "load.html", {
        "import_id": import_id,
        "many": [i for i in shown if i.kind == "many"],
        "windows": [i for i in shown if i.kind == "window"],
        "max_lessons": max_lessons,
        "window_hours": int(window_hours) if float(window_hours).is_integer() else window_hours,
        "day": chosen,
        "day_counts": {d: sum(1 for i in items if i.day == d) for d in range(len(DAY_NAMES_RU))},
        "day_total": len(items),
        "day_names": DAY_NAMES_RU, "day_short": DAY_SHORT_RU,
    })


@app.get("/report/{import_id}/free-rooms", response_class=HTMLResponse)
def free_rooms_page(request: Request, import_id: str, day: str = ""):
    lessons, notes = _load_import(import_id)
    if lessons is None:
        return HTMLResponse("<h1>Импорт не найден</h1><p><a href='/'>На главную</a></p>", status_code=404)

    rooms = cabinets_module.department_rooms(cabinets_module.get_rooms(cabinets_module.load_config(CABINETS_PATH)))
    days = sorted(timegrid.build_day_grids(lessons).keys())
    chosen = _parse_day(day)
    if chosen not in days:
        chosen = days[0] if days else 0
    view = cabinets_module.free_rooms_day(lessons, chosen, rooms, (notes or {}).get("special_slots", []))
    return templates.TemplateResponse(request, "free_rooms.html", {
        "import_id": import_id,
        "days": days,
        "day": chosen,
        "day_names": DAY_NAMES_RU,
        "view": view,
        "has_rooms": bool(rooms),
        "room_count": len(rooms),
        "room_label": _room_label,
    })


def _involves_accompanist(c) -> bool:
    return bool(c.lesson_a.accompanist_name or c.lesson_b.accompanist_name)


def _lesson_names(l) -> list[str]:
    return [n for n in (l.teacher_name, l.accompanist_name, l.student_name) if n]


def _person_match_key(name: str) -> str:
    """Ключ сравнения ФИО для фильтра: без учёного звания и регистра — та же
    нормализация, что и в conflicts.py, чтобы 'доц. Долгачева С.А.' и
    'Долгачева С.А.' считались одним человеком (см. strip_academic_title)."""
    return strip_academic_title(name).casefold()


def _conflict_matches_person(c, needle_key: str) -> bool:
    return any(
        _person_match_key(name) == needle_key
        for name in _lesson_names(c.lesson_a) + _lesson_names(c.lesson_b)
    )


def _filter_conflicts_by_person(conflicts: list, person: str) -> list:
    """Оставляет только накладки, где ФИО препода/концертмейстера/студента
    (с любой из двух сторон накладки) совпадает с person без учёта учёного
    звания и регистра — выбор идёт из выпадающего списка готовых ФИО, поэтому
    сравниваем не подстрокой, а целиком. Пустой person означает "без фильтра"."""
    needle_key = _person_match_key(person.strip())
    if not needle_key:
        return conflicts
    return [c for c in conflicts if _conflict_matches_person(c, needle_key)]


def _collect_all_names(lessons: list) -> list[str]:
    """Все встречающиеся ФИО (препод./концертмейстер/студент) — для выпадающего
    списка на странице отчёта. 'доц. Х' и 'Х' — один и тот же человек (см.
    strip_academic_title), схлопываем в одну запись в списке; предпочитаем
    форму без звания как более узнаваемую методисту."""
    by_key: dict[str, str] = {}
    for l in lessons:
        for name in _lesson_names(l):
            key = _person_match_key(name)
            current = by_key.get(key)
            if current is None or (strip_academic_title(current) != current and strip_academic_title(name) == name):
                by_key[key] = name
    return sorted(by_key.values(), key=str.casefold)


DAY_SHORT_RU = ["Пн", "Вт", "Ср", "Чт", "Пт", "Сб"]



def _parse_day(day: str) -> int | None:
    try:
        d = int(day)
    except (TypeError, ValueError):
        return None
    return d if 0 <= d < len(DAY_NAMES_RU) else None


def _conflict_in_room(c, room: str) -> bool:
    return room in (c.lesson_a.room_normalized, c.lesson_b.room_normalized)


def _room_label(key: str) -> str:
    """'421' -> '421'; '1:405' (чужой корпус) -> 'корп. 1, ауд. 405'."""
    building, sep, num = key.partition(":")
    return f"корп. {building}, ауд. {num}" if sep else key


def _fmt_range(l) -> str:
    end = l.end_minutes
    return f"{l.start_time.strftime('%H:%M')}–{end // 60:02d}:{end % 60:02d}"


def _overlap_range(c) -> str:
    """Время, когда два занятия реально накладываются: 'с 14:35 до 15:20'."""
    start = max(c.lesson_a.start_minutes, c.lesson_b.start_minutes)
    end = min(c.lesson_a.end_minutes, c.lesson_b.end_minutes)
    return f"{start // 60:02d}:{start % 60:02d}–{end // 60:02d}:{end % 60:02d}"


# Браузер кэширует style.css; после обновления программы (и правок стилей) без
# метки версии у методиста могла бы остаться старая вёрстка.
templates.env.globals["app_version"] = APP_VERSION
templates.env.globals["cached_update"] = cached_update
# Вкладка «Свободные кабинеты» появляется только когда методист указал кабинеты своей кафедры.
templates.env.globals["has_department_rooms"] = lambda: bool(
    cabinets_module.get_rooms(cabinets_module.load_config(CABINETS_PATH))
)
templates.env.globals["changelog_entries"] = changelog.entries_for_display
templates.env.filters["md"] = render_markdown_lite
templates.env.globals["css_stamp"] = int(os.path.getmtime(os.path.join(APP_DIR, "static", "style.css")))
templates.env.globals["overlap_range"] = _overlap_range
templates.env.globals["conflict_key"] = conflict_key
templates.env.globals["fmt_range"] = _fmt_range
templates.env.globals["day_short"] = DAY_SHORT_RU


class _Filters:
    """Параметры отчёта из адресной строки; одни и те же для страницы,
    выгрузок в Excel/Word и JSON-API — чтобы файл содержал ровно то, что
    методист видит на экране."""

    def __init__(self, hide_accompanist: bool = False, person: str = "", day: str = "",
                 room: str = "", show_dismissed: bool = False, changes: str = "", compare: str = ""):
        self.changes = "new" if changes == "new" else ""     # "new" — только новые с прошлой проверки
        self.compare = compare.strip()                        # "" — прошлая, "off" — не сравнивать, иначе id
        self.hide_accompanist = hide_accompanist
        self.person = person.strip()
        self.day = _parse_day(day)
        self.room = room.strip()
        self.show_dismissed = show_dismissed

    def as_params(self, **override) -> dict:
        p = {
            "hide_accompanist": "true" if self.hide_accompanist else "",
            "person": self.person,
            "day": "" if self.day is None else str(self.day),
            "room": self.room,
            "show_dismissed": "true" if self.show_dismissed else "",
            "changes": self.changes,
            "compare": self.compare,
        }
        p.update(override)
        return {k: v for k, v in p.items() if v}


class _Selection:
    def __init__(self, conflicts: list, pairings: list, dismissed: list, day_counts: dict,
                 day_options: list, room_options: list, dismissed_total: int):
        self.conflicts = conflicts          # активные (не помеченные «не накладка»), без пар с концертмейстером
        self.pairings = pairings            # пары «преподаватель + концертмейстер» для сверки (не накладки)
        self.dismissed = dismissed          # помеченные, но подходящие под фильтры
        self.day_counts = day_counts        # {день: число активных} при остальных фильтрах
        self.day_options = day_options
        self.room_options = room_options
        self.dismissed_total = dismissed_total


_room_sort_key = cabinets_module.room_sort_key


def _select_conflicts(lessons: list, f: _Filters, new_keys: set | None = None) -> _Selection:
    all_conflicts = find_conflicts(lessons)
    dismissed_keys = dismissed_module.load_dismissed(DISMISSED_PATH)

    base = all_conflicts
    if f.hide_accompanist:
        base = [c for c in base if not _involves_accompanist(c)]
    base = _filter_conflicts_by_person(base, f.person)
    if f.room:
        base = [c for c in base if _conflict_in_room(c, f.room)]
    if f.changes == "new" and new_keys is not None:
        base = [c for c in base if conflict_key(c) in new_keys]

    def is_dismissed(c) -> bool:
        return conflict_key(c) in dismissed_keys

    day_counts: dict[int, int] = {}
    for c in base:
        if c.type != ConflictType.ACCOMPANIST_PAIRING and not is_dismissed(c):
            day_counts[c.day_of_week] = day_counts.get(c.day_of_week, 0) + 1

    picked = [c for c in base if f.day is None or c.day_of_week == f.day]
    is_pairing = lambda c: c.type == ConflictType.ACCOMPANIST_PAIRING
    active = [c for c in picked if not is_pairing(c) and not is_dismissed(c)]
    pairings = [c for c in picked if is_pairing(c) and not is_dismissed(c)]
    dismissed = [c for c in picked if is_dismissed(c)]

    rooms = {r for c in all_conflicts for r in (c.lesson_a.room_normalized, c.lesson_b.room_normalized) if r}
    if f.room:
        rooms.add(f.room)
    room_options = [(r, _room_label(r)) for r in sorted(rooms, key=_room_sort_key)]

    return _Selection(
        active, pairings, dismissed, day_counts, sorted({c.day_of_week for c in all_conflicts}),
        room_options, sum(1 for c in all_conflicts if is_dismissed(c)),
    )


def _load_import(import_id: str):
    conn = db.get_connection(DB_PATH)
    try:
        if not db.import_exists(conn, import_id):
            return None, None
        return db.load_lessons(conn, import_id), db.load_notes(conn, import_id)
    finally:
        conn.close()


def _comparison(import_id: str, lessons: list, f: _Filters):
    """Сравнение с прошлой (или выбранной) проверкой — compare.Comparison или None.
    Вторым значением — список других проверок для выпадающего списка."""
    conn = db.get_connection(DB_PATH)
    try:
        rows = db.list_imports(conn)
        prev = compare_module.pick_previous(rows, import_id, f.compare)
        options = [(r["id"], _format_ru_datetime(r["created_at"])) for r in rows if r["id"] != import_id]
        if prev is None:
            return None, options
        prev_lessons = db.load_lessons(conn, prev["id"])
        prev_created = prev["created_at"]
    finally:
        conn.close()
    cmp = compare_module.compare(
        find_conflicts(lessons), find_conflicts(prev_lessons),
        prev["id"], prev_created, len(prev_lessons), len(lessons),
    )
    return cmp, options


def _new_keys_for_export(import_id: str, lessons: list, f: _Filters):
    """Для выгрузок: ключи новых накладок, если на экране включён «только новые»."""
    if f.changes != "new":
        return None
    cmp, _ = _comparison(import_id, lessons, f)
    return cmp.new_keys if cmp else None


def _build_report_context(import_id: str, f: _Filters | None = None) -> dict | None:
    f = f or _Filters()
    lessons, notes = _load_import(import_id)
    if lessons is None:
        return None
    cmp, compare_options = _comparison(import_id, lessons, f)
    new_keys = cmp.new_keys if cmp else None
    sel = _select_conflicts(lessons, f, new_keys)

    def grouped(conflicts: list) -> list[dict]:
        by_type: dict[str, list] = {}
        for c in conflicts:
            by_type.setdefault(c.type.value, []).append(c)
        return [
            {"type": ctype, "label": CONFLICT_LABELS[ctype], "conflicts": by_type.get(ctype.value, [])}
            for ctype in ConflictType
        ]

    individual_count = sum(1 for l in lessons if l.lesson_type.value == "individual")
    irregular_times = [
        {
            "who": i.lesson.teacher_name or i.lesson.accompanist_name or "—",
            "student": i.lesson.student_name,
            "room": i.lesson.room_raw,
            "text": timegrid.describe(i),
            "source": i.lesson.source.label(),
        }
        for i in timegrid.find_irregular_times(lessons)
    ]
    lost_notes = [n for n in notes.get("dropped_sheet_notes", []) if "ВНИМАНИЕ" in n]

    def link(**override) -> str:
        params = f.as_params(**override)
        return f"/report/{import_id}" + (f"?{urlencode(params)}" if params else "")

    export_params = f.as_params(show_dismissed="")
    return {
        "import_id": import_id,
        "notes": notes,
        "lost_notes": lost_notes,
        "irregular_times": irregular_times,
        "lesson_count": len(lessons),
        "individual_count": individual_count,
        "group_count": len(lessons) - individual_count,
        "certain_count": sum(1 for c in sel.conflicts if c.is_certain),
        "review_count": sum(1 for c in sel.conflicts if not c.is_certain),
        "dismissed_total": sel.dismissed_total,
        "f": f,
        "hide_accompanist": f.hide_accompanist,
        "person": f.person,
        "day": f.day,
        "room": f.room,
        "show_dismissed": f.show_dismissed,
        "all_names": _collect_all_names(lessons),
        "day_options": sel.day_options,
        "day_counts": sel.day_counts,
        "day_total": sum(sel.day_counts.values()),
        "room_options": sel.room_options,
        "conflict_groups": [g for g in grouped(sel.conflicts) if g["type"] != ConflictType.ACCOMPANIST_PAIRING],
        "pairing_groups": [g for g in grouped(sel.pairings) if g["type"] == ConflictType.ACCOMPANIST_PAIRING],
        "pairing_count": len(sel.pairings),
        "pairing_review_count": sum(1 for c in sel.pairings if not c.is_certain),
        "dismissed_groups": grouped(sel.dismissed),
        "dismissed_shown": len(sel.dismissed),
        "cmp": cmp,
        "cmp_label": _format_ru_datetime(cmp.prev_created) if cmp else "",
        "cmp_fixed": [
            {
                "day": DAY_SHORT_RU[c.day_of_week], "when": _overlap_range(c), "kind": CONFLICT_LABELS[c.type],
                "a": c.lesson_a.teacher_name or c.lesson_a.accompanist_name or "—",
                "b": c.lesson_b.teacher_name or c.lesson_b.accompanist_name or "—",
                "room": c.lesson_a.room_raw or c.lesson_b.room_raw or "",
            }
            for c in (cmp.fixed if cmp else [])
        ],
        "compare_options": compare_options,
        "new_keys": new_keys or set(),
        "only_new": f.changes == "new" and cmp is not None,
        "day_names": DAY_NAMES_RU,
        "qs": urlencode(f.as_params()),
        "link": link,
        "export_qs": f"?{urlencode(export_params)}" if export_params else "",
        "has_filters": bool(f.hide_accompanist or f.person or f.day is not None or f.room or f.changes),
    }


@app.get("/report/{import_id}", response_class=HTMLResponse)
def report(request: Request, import_id: str, hide_accompanist: bool = False, person: str = "",
           day: str = "", room: str = "", show_dismissed: bool = False, changes: str = "", compare: str = ""):
    ctx = _build_report_context(import_id, _Filters(hide_accompanist, person, day, room, show_dismissed, changes, compare))
    if ctx is None:
        return HTMLResponse("<h1>Импорт не найден</h1><p><a href='/'>На главную</a></p>", status_code=404)
    return templates.TemplateResponse(request, "report.html", ctx)


_KEY_RE = re.compile(r"^[0-9a-f]{16}$")


def _back_to_report(import_id: str, qs: str) -> RedirectResponse:
    return RedirectResponse(url=f"/report/{import_id}" + (f"?{qs}" if qs else ""), status_code=303)


def _wants_json(request: Request) -> bool:
    return request.headers.get("x-requested-with") == "fetch"


@app.post("/report/{import_id}/dismiss")
def dismiss_conflict(request: Request, import_id: str, key: str = Form(...), label: str = Form(""), qs: str = Form("")):
    if _KEY_RE.match(key):
        dismissed_module.dismiss(DISMISSED_PATH, key, label[:300])
    if _wants_json(request):
        return JSONResponse({"ok": True})
    return _back_to_report(import_id, qs)


@app.post("/report/{import_id}/restore")
def restore_conflict(request: Request, import_id: str, key: str = Form(""), qs: str = Form("")):
    if key == "all":
        dismissed_module.restore_all(DISMISSED_PATH)
    elif _KEY_RE.match(key):
        dismissed_module.restore(DISMISSED_PATH, key)
    if _wants_json(request):
        return JSONResponse({"ok": True})
    return _back_to_report(import_id, qs)


@app.get("/api/report/{import_id}")
def api_report(import_id: str, hide_accompanist: bool = False, person: str = "", day: str = "", room: str = "",
               changes: str = "", compare: str = ""):
    lessons, _ = _load_import(import_id)
    if lessons is None:
        return {"error": "import not found"}
    f = _Filters(hide_accompanist, person, day, room, changes=changes, compare=compare)
    sel = _select_conflicts(lessons, f, _new_keys_for_export(import_id, lessons, f))

    def lesson_json(l):
        return {
            "type": l.lesson_type.value, "day_of_week": l.day_of_week,
            "start_time": l.start_time.strftime("%H:%M"), "duration_minutes": l.duration_minutes,
            "teacher_name": l.teacher_name, "accompanist_name": l.accompanist_name,
            "student_name": l.student_name, "group": l.group_raw, "subject": l.subject,
            "room": l.room_raw, "source": l.source.label(),
        }

    return {
        "import_id": import_id,
        "lesson_count": len(lessons),
        "conflicts": [
            {
                "type": c.type.value, "day_of_week": c.day_of_week, "is_certain": c.is_certain,
                "note": c.note, "a": lesson_json(c.lesson_a), "b": lesson_json(c.lesson_b),
            }
            for c in sel.conflicts
        ],
    }


@app.get("/report/{import_id}/export.xlsx")
def export_xlsx(import_id: str, hide_accompanist: bool = False, person: str = "", day: str = "", room: str = "",
                changes: str = "", compare: str = ""):
    import openpyxl
    from openpyxl.utils import get_column_letter

    lessons, _ = _load_import(import_id)
    if lessons is None:
        return HTMLResponse("Импорт не найден", status_code=404)
    f = _Filters(hide_accompanist, person, day, room, changes=changes, compare=compare)
    sel = _select_conflicts(lessons, f, _new_keys_for_export(import_id, lessons, f))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Накладки"
    headers = [
        "Тип накладки", "Требует проверки", "День", "Примечание",
        "A: время", "A: кто", "A: студент", "A: группа", "A: предмет", "A: аудитория", "A: источник",
        "B: время", "B: кто", "B: студент", "B: группа", "B: предмет", "B: аудитория", "B: источник",
    ]

    def who(l):
        return l.teacher_name or l.accompanist_name or ""

    def fill(sheet, conflicts):
        sheet.append(headers)
        for c in conflicts:
            a, b = c.lesson_a, c.lesson_b
            sheet.append([
                CONFLICT_LABELS[c.type], "да" if not c.is_certain else "", DAY_NAMES_RU[c.day_of_week], c.note or "",
                _fmt_range(a), who(a), a.student_name or "", a.group_raw or "", a.subject or "", a.room_raw or "", a.source.label(),
                _fmt_range(b), who(b), b.student_name or "", b.group_raw or "", b.subject or "", b.room_raw or "", b.source.label(),
            ])
        for i, _ in enumerate(headers, start=1):
            sheet.column_dimensions[get_column_letter(i)].width = 20

    fill(ws, sel.conflicts)
    if sel.pairings:
        fill(wb.create_sheet("Проверить концертмейстера"), sel.pairings)

    import io
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=nakladki_{import_id}.xlsx"},
    )


def _lesson_line(l) -> str:
    """Одна строка с описанием занятия — для отчёта в Word (там нет места
    под многоколоночную вёрстку как в HTML/Excel, зато читается как обычный текст)."""
    person = l.teacher_name or l.accompanist_name or "—"
    bits = [person, l.start_time.strftime("%H:%M")]
    if l.student_name:
        bits.append(f"студент {l.student_name}")
    bits.append(f"гр.{l.group_raw or '—'}")
    bits.append(l.subject or "—")
    bits.append(f"ауд.{l.room_raw or '—'}")
    return " · ".join(bits) + f" ({l.source.label()})"


def _build_conflict_docx(ctx: dict):
    """Собирает документ Word с тем же содержимым, что и страница отчёта —
    чтобы методисту не приходилось вручную копировать страницу и вставлять
    в Word."""
    from docx import Document

    doc = Document()
    doc.add_heading("Отчёт по накладкам", level=1)

    summary_bits = [
        f"Занятий всего: {ctx['lesson_count']} ({ctx['individual_count']} индивид. + {ctx['group_count']} групп.)",
        f"Явных накладок: {ctx['certain_count']}",
        f"Требуют ручной проверки: {ctx['review_count']}",
    ]
    if ctx["person"]:
        summary_bits.append(f"Показаны только накладки: {ctx['person']}")
    if ctx["day"] is not None:
        summary_bits.append(f"День: {DAY_NAMES_RU[ctx['day']]}")
    if ctx["room"]:
        summary_bits.append(f"Кабинет: {_room_label(ctx['room'])}")
    if ctx["hide_accompanist"]:
        summary_bits.append("Накладки с участием концертмейстера скрыты")
    doc.add_paragraph(" · ".join(summary_bits))

    if ctx["has_filters"] and ctx["certain_count"] == 0 and ctx["review_count"] == 0:
        doc.add_paragraph("По выбранным условиям накладок не найдено — всё в порядке.")

    day_names = ctx["day_names"]
    for group in ctx["conflict_groups"]:
        if not group["conflicts"]:
            continue
        doc.add_heading(f"{group['label']} ({len(group['conflicts'])})", level=2)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "День", "Запись A", "Запись B", "Примечание"
        for c in group["conflicts"]:
            row = table.add_row().cells
            row[0].text = day_names[c.day_of_week]
            row[1].text = _lesson_line(c.lesson_a)
            row[2].text = _lesson_line(c.lesson_b)
            note = ("[Требует проверки] " if not c.is_certain else "") + (c.note or "")
            row[3].text = note.strip()

    for group in ctx["pairing_groups"]:
        if not group["conflicts"]:
            continue
        doc.add_heading(f"{group['label']} ({len(group['conflicts'])})", level=2)
        doc.add_paragraph(
            "Преподаватель и концертмейстер у одного студента в одном кабинете одновременно — "
            "обычная практика, это не накладки. Сверьте, что концертмейстер назначен верно."
        )
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "День", "Преподаватель", "Концертмейстер", "Примечание"
        for c in group["conflicts"]:
            row = table.add_row().cells
            row[0].text = day_names[c.day_of_week]
            row[1].text = _lesson_line(c.lesson_a)
            row[2].text = _lesson_line(c.lesson_b)
            row[3].text = ("[Требует проверки] " if not c.is_certain else "") + (c.note or "")

    return doc


@app.get("/report/{import_id}/export.docx")
def export_docx(import_id: str, hide_accompanist: bool = False, person: str = "", day: str = "", room: str = "",
                changes: str = "", compare: str = ""):
    ctx = _build_report_context(import_id, _Filters(hide_accompanist, person, day, room, changes=changes, compare=compare))
    if ctx is None:
        return HTMLResponse("Импорт не найден", status_code=404)

    doc = _build_conflict_docx(ctx)

    import io
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename=otchet_nakladki_{import_id}.docx"},
    )
